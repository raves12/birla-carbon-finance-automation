"""Small python-docx helpers shared by the BCFA deliverable builders."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x38, 0x64)
TEAL = RGBColor(0x0B, 0x6E, 0x69)
GREY = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = '1F3864'
ALT_FILL = 'EEF2F8'
NOTE_FILL = 'FFF4E5'
SAY_FILL = 'E8F4F2'


def new_document(title, subtitle, meta_rows, landscape=False):
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    st.paragraph_format.space_after = Pt(4)
    for lvl, size in ((1, 16), (2, 13), (3, 11.5)):
        h = doc.styles['Heading %d' % lvl]
        h.font.name = 'Calibri'
        h.font.size = Pt(size)
        h.font.color.rgb = NAVY
        h.font.bold = True
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        h.paragraph_format.space_after = Pt(4)
    # Cover
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(14)
    r.font.color.rgb = TEAL
    doc.add_paragraph()
    t = table(doc, ['Item', 'Value'], meta_rows, widths=[5, 11])
    _ = t
    footer_page_numbers(doc, title)
    page_break(doc)
    return doc


def footer_page_numbers(doc, text):
    for s in doc.sections:
        p = s.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text + '  |  Page ')
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
        _field(p, 'PAGE', 8)


def _field(paragraph, code, size=None):
    run = paragraph.add_run()
    if size:
        run.font.size = Pt(size)
    b = OxmlElement('w:fldChar')
    b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText')
    i.set(qn('xml:space'), 'preserve')
    i.text = code
    s = OxmlElement('w:fldChar')
    s.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t')
    t.text = '1' if code == 'PAGE' else 'Right-click here and choose "Update Field" to build the table of contents.'
    e = OxmlElement('w:fldChar')
    e.set(qn('w:fldCharType'), 'end')
    run._r.append(b)
    run._r.append(i)
    run._r.append(s)
    run._r.append(t)
    run._r.append(e)


def toc(doc, levels='1-2'):
    t = doc.add_paragraph()
    r = t.add_run('Contents')
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "%s" \\h \\z \\u' % levels)
    page_break(doc)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def h(doc, text, level=1):
    return doc.add_heading(text, level=level)


def para(doc, text, bold_lead=None, italic=False, size=None, color=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def bullets(doc, items, style='List Bullet'):
    for it in items:
        p = doc.add_paragraph(style=style)
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)


def numbered(doc, items):
    bullets(doc, items, style='List Number')


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)


def table(doc, headers, rows, widths=None, font=8.5, zebra=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0]
    _repeat_header(hdr)
    for i, text in enumerate(headers):
        c = hdr.cells[i]
        c.text = ''
        r = c.paragraphs[0].add_run(str(text))
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, HEADER_FILL)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ''
            lines = str('' if val is None else val).split('\n')
            p = c.paragraphs[0]
            for k, ln in enumerate(lines):
                if k:
                    p = c.add_paragraph()
                r = p.add_run(ln)
                r.font.size = Pt(font)
                p.paragraph_format.space_after = Pt(0)
            if zebra and n % 2 == 1:
                _shade(c, ALT_FILL)
    if widths:
        set_widths(t, widths)
    doc.add_paragraph()
    return t


def set_widths(t, widths):
    """Fixed column widths that Word honours: fixed layout, grid columns and every cell."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        if i < len(widths):
            gc.set(qn('w:w'), str(int(Cm(widths[i]).twips)))
    for row in t.rows:
        for i, w in enumerate(widths):
            if i < len(row.cells):
                row.cells[i].width = Cm(w)


def box(doc, title, lines, fill=NOTE_FILL, title_color=None):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]
    _shade(c, fill)
    c.text = ''
    p = c.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = title_color or NAVY
    for ln in lines:
        p = c.add_paragraph()
        if isinstance(ln, tuple):
            r = p.add_run(ln[0])
            r.bold = True
            p.add_run(ln[1])
        else:
            p.add_run(ln)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph()
    return t


def code(doc, text, size=8):
    for ln in text.split('\n'):
        p = doc.add_paragraph()
        r = p.add_run(ln if ln else ' ')
        r.font.name = 'Consolas'
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
    doc.add_paragraph()
